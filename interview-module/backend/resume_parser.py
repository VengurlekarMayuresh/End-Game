"""
Resume and Job Description Parser Module for HireSense AI.
Supports PDF (.pdf), Microsoft Word (.docx), and plain text (.txt) file parsing.
"""

import io
import re
from pathlib import Path
from pypdf import PdfReader
from docx import Document

from nlp_extract import extract_keywords


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """
    Extracts raw text from PDF, DOCX, or TXT file bytes.
    """
    ext = Path(filename).suffix.lower()
    
    if ext == ".pdf":
        return _extract_from_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return _extract_from_docx(file_bytes)
    elif ext in (".txt", ".md"):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore")
    else:
        raise ValueError(f"Unsupported file format: '{ext}'. Please upload a .pdf, .docx, or .txt file.")


def _extract_from_pdf(file_bytes: bytes) -> str:
    """Extracts text page by page from PDF."""
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n\n".join(text_parts).strip()


def _extract_from_docx(file_bytes: bytes) -> str:
    """Extracts text from paragraphs and tables in DOCX."""
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                text_parts.append(row_text)
    return "\n".join(text_parts).strip()


def parse_resume_data(raw_text: str) -> dict:
    """
    Parses resume text into structured data dictionary with skills, projects,
    and metadata using NLP extraction and heuristics.
    """
    # 1. NLP Keyword extraction for high-precision technical skills
    kw_result = extract_keywords(raw_text)
    skills = kw_result.get("matched_skills", [])

    # 2. Extract Candidate Name (heuristics: first non-empty line or title)
    name = _extract_name(raw_text)

    # 3. Extract Projects
    projects = _extract_projects(raw_text, skills)

    # 4. Extract Education
    education = _extract_education(raw_text)

    return {
        "name": name,
        "raw_text": raw_text,
        "skills": skills,
        "projects": projects,
        "education": education,
    }


def parse_job_description(jd_text: str, role: str = None) -> dict:
    """
    Parses raw JD text to detect required skills, target role, and core topics.
    """
    kw_result = extract_keywords(jd_text)
    required_skills = kw_result.get("matched_skills", [])

    # Detect role if not provided
    if not role:
        role_match = re.search(r"(?:role|position|job title|looking for a|seeking a)\s*:?\s*([^\n\r,\.]+)", jd_text, re.IGNORECASE)
        role = role_match.group(1).strip() if role_match else "Software Engineer"
        if len(role) > 50:
            role = role[:50].strip()

    # Core topics selection based on skills in JD
    core_topics = []
    text_lower = jd_text.lower()
    if any(k in text_lower for k in ["dsa", "data structures", "algorithms", "problem solving"]):
        core_topics.append("DSA")
    if any(k in text_lower for k in ["os", "operating system", "linux", "concurrency", "thread"]):
        core_topics.append("OS")
    if any(k in text_lower for k in ["network", "tcp", "udp", "http", "socket"]):
        core_topics.append("Computer Networks")
    if any(k in text_lower for k in ["sql", "database", "postgres", "mysql", "mongodb", "dbms"]):
        core_topics.append("DBMS")

    # Fallback to standard CS subjects if none explicitly flagged
    if not core_topics:
        core_topics = ["DSA", "OS", "Computer Networks"]

    return {
        "role": role,
        "required_skills": required_skills,
        "core_topics": core_topics,
        "raw_text": jd_text,
    }


def _extract_name(text: str) -> str:
    """Heuristic extraction of candidate name."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines:
        first_line = lines[0]
        # Ignore if it starts with 'Resume' or 'Curriculum Vitae'
        if not re.search(r"^(resume|curriculum|cv|profile)", first_line, re.I):
            clean_name = re.sub(r"[^A-Za-z\s\.\-]", "", first_line).strip()
            if 2 <= len(clean_name.split()) <= 4:
                return clean_name
    return "Candidate"


def _extract_projects(text: str, detected_skills: list[str]) -> list[dict]:
    """Extracts project entries from project sections in resume."""
    projects = []
    
    # Try to find a Projects section
    pattern = re.compile(r"(?:PROJECTS|ACADEMIC PROJECTS|PERSONAL PROJECTS|KEY PROJECTS)\s*\n(.*?)(?=\n[A-Z\s]{4,20}\n|\Z)", re.DOTALL | re.IGNORECASE)
    match = pattern.search(text)
    
    if match:
        section_text = match.group(1).strip()
        # Split projects by bullet points or empty lines
        entries = re.split(r"\n(?=[•\-\*A-Z0-9][^\n]+(?::|-|—))|\n\n+", section_text)
        for entry in entries:
            entry_clean = entry.strip()
            if len(entry_clean) > 20:
                lines = entry_clean.splitlines()
                title_line = lines[0].strip("•-* ")
                title_match = re.split(r"[:\-|—]", title_line)
                proj_name = title_match[0].strip() if title_match else "Project"
                
                # Associated tech stack
                entry_skills = [s for s in detected_skills if s.lower() in entry_clean.lower()]
                tech_stack = ", ".join(entry_skills) if entry_skills else "Full Stack"
                
                projects.append({
                    "name": proj_name,
                    "description": entry_clean[:300],
                    "tech_stack": tech_stack
                })
    
    # Fallback if no explicit project section was found
    if not projects:
        projects.append({
            "name": "Featured Project",
            "description": "Technical projects and software systems built by the candidate.",
            "tech_stack": ", ".join(detected_skills[:4]) if detected_skills else "General Tech Stack"
        })

    return projects[:3]


def _extract_education(text: str) -> list[dict]:
    """Extracts education mentions (Degree, University)."""
    education = []
    degree_matches = re.findall(r"(B\.?Tech|B\.?E\.?|B\.?S\.?|M\.?Tech|M\.?S\.?|Master|Bachelor|Diploma|Computer Science|Engineering)[A-Za-z\s,]+", text, re.I)
    for m in degree_matches[:2]:
        education.append({"degree": m.strip()})
    return education
